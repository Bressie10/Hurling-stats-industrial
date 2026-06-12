from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


OUT_DIR = Path("pitch")
SLIDE_DIR = OUT_DIR / "rendered_slides"
PPTX_PATH = OUT_DIR / "GAAstat_90_second_pitch.pptx"
VOICE_PPTX_PATH = OUT_DIR / "GAAstat_voice_first_90_second_pitch.pptx"
SCRIPT_PATH = OUT_DIR / "GAAstat_voice_first_90_second_pitch_script.md"
AUDIT_PATH = OUT_DIR / "GAAstat_current_app_audit.md"

W, H = 1920, 1080

COL = {
    "bg": "#07110F",
    "panel": "#10241E",
    "panel2": "#15352B",
    "stroke": "#2B5A4B",
    "lime": "#A8E63D",
    "lime2": "#D6FF68",
    "orange": "#FF8C00",
    "white": "#F5FFF7",
    "muted": "#B5C8BD",
    "dim": "#728279",
    "red": "#F05E59",
    "amber": "#F5B041",
    "blue": "#45B5FF",
    "pitch": "#226B3F",
}

FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
FONT_BLACK = "/System/Library/Fonts/Supplemental/Arial Black.ttf"


def font(size, bold=False, black=False):
    path = FONT_BLACK if black else FONT_BOLD if bold else FONT
    return ImageFont.truetype(path, size)


def text_size(draw, text, fnt):
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0], box[3] - box[1]


def wrap(draw, text, fnt, max_w):
    words = text.split()
    lines = []
    line = ""
    for word in words:
        test = word if not line else f"{line} {word}"
        if text_size(draw, test, fnt)[0] <= max_w:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_text(draw, text, xy, size, color=COL["white"], bold=False, black=False, max_w=None,
              line_gap=8, anchor="la", align="left"):
    fnt = font(size, bold=bold, black=black)
    x, y = xy
    lines = text.split("\n") if "\n" in text else wrap(draw, text, fnt, max_w) if max_w else [text]
    line_h = size + line_gap
    for i, line in enumerate(lines):
      line_x = x
      if align == "center" and max_w:
          line_x = x + (max_w - text_size(draw, line, fnt)[0]) // 2
      elif align == "right" and max_w:
          line_x = x + max_w - text_size(draw, line, fnt)[0]
      draw.text((line_x, y + i * line_h), line, font=fnt, fill=color, anchor=anchor)
    return y + len(lines) * line_h


def rounded(draw, box, fill, outline=None, width=2, r=24):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)


def base_slide():
    img = Image.new("RGB", (W, H), COL["bg"])
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, W, 16), fill=COL["lime"])
    draw.rectangle((0, H - 16, W, H), fill=COL["orange"])
    draw.rectangle((W - 42, 0, W, H), fill="#0D271F")
    # Very subtle field-like lines.
    for x in range(-120, W, 260):
        draw.line((x, 160, x + 360, H - 90), fill="#12372D", width=2)
    return img, draw


def logo(draw, x=92, y=76):
    draw.ellipse((x, y, x + 66, y + 66), fill=COL["lime"])
    draw_text(draw, "G", (x + 19, y + 10), 36, COL["bg"], bold=True)
    draw_text(draw, "GAA", (x + 90, y + 12), 34, COL["lime"], black=True)
    draw_text(draw, "stat", (x + 178, y + 12), 34, COL["white"], bold=True)


def header(draw, label):
    logo(draw)
    draw_text(draw, label.upper(), (W - 360, 88), 18, COL["dim"], bold=True, max_w=270, align="right")
    draw.line((92, 158, W - 120, 158), fill=COL["stroke"], width=2)


def pill(draw, text, box, fill, text_color=COL["bg"], size=20):
    rounded(draw, box, fill, r=13)
    x1, y1, x2, y2 = box
    tw, th = text_size(draw, text, font(size, bold=True))
    draw.text((x1 + (x2 - x1 - tw) / 2, y1 + (y2 - y1 - th) / 2 - 2), text, font=font(size, bold=True), fill=text_color)


def title_block(draw, eyebrow, title, subtitle=None):
    draw_text(draw, eyebrow.upper(), (100, 218), 20, COL["orange"], bold=True)
    draw_text(draw, title, (100, 270), 64, COL["white"], black=True, max_w=820, line_gap=14)
    if subtitle:
        draw_text(draw, subtitle, (104, 440), 30, COL["muted"], bold=True, max_w=760, line_gap=8)


def phone(draw, x, y, scale=1.0):
    w, h = int(420 * scale), int(760 * scale)
    rounded(draw, (x, y, x + w, y + h), "#09110F", outline="#4B5E56", width=3, r=int(38 * scale))
    rounded(draw, (x + 24, y + 36, x + w - 24, y + h - 36), "#092118", outline=COL["stroke"], width=2, r=int(24 * scale))
    draw.rounded_rectangle((x + w // 2 - 40, y + 18, x + w // 2 + 40, y + 24), radius=3, fill="#4B5E56")
    sx = x + 48
    sw = w - 96
    draw_text(draw, "Championship vs Ballymore", (sx, y + 70), int(16 * scale), COL["muted"], bold=True)
    draw_text(draw, "2-14", (sx + 86, y + 120), int(38 * scale), COL["lime"], black=True)
    draw_text(draw, "-", (sx + 178, y + 124), int(34 * scale), COL["dim"], bold=True)
    draw_text(draw, "1-11", (sx + 214, y + 120), int(38 * scale), COL["white"], black=True)
    rounded(draw, (sx, y + 200, sx + sw, y + 262), COL["panel2"], outline=COL["stroke"], width=2, r=12)
    draw_text(draw, "Sideline AI", (sx + 24, y + 218), int(17 * scale), COL["white"], bold=True)
    pill(draw, "READY", (sx + sw - 92, y + 210, sx + sw - 20, y + 250), COL["lime"], size=int(12 * scale))
    rounded(draw, (sx + 28, y + 302, sx + sw - 28, y + 388), COL["lime"], r=12)
    draw_text(draw, "Hold to talk", (sx + 92, y + 326), int(22 * scale), COL["bg"], bold=True)
    messages = [
        ("YOU", "Point for number 11", COL["blue"], COL["panel2"]),
        ("AI", "Pending point #11. Confirm?", COL["lime"], COL["panel"]),
    ]
    yy = y + 432
    for label, msg, accent, fill in messages:
        rounded(draw, (sx, yy, sx + sw, yy + 78), fill, outline=COL["stroke"], width=2, r=13)
        draw_text(draw, label, (sx + 24, yy + 15), int(13 * scale), accent, bold=True)
        draw_text(draw, msg, (sx + 24, yy + 42), int(18 * scale), COL["white"], bold=True)
        yy += 96


def card(draw, box, title, body, accent=COL["lime"]):
    rounded(draw, box, COL["panel"], outline=COL["stroke"], width=2, r=22)
    x1, y1, x2, y2 = box
    draw.ellipse((x1 + 28, y1 + 28, x1 + 74, y1 + 74), fill=accent)
    draw_text(draw, title, (x1 + 30, y1 + 110), 34, COL["white"], black=True)
    draw_text(draw, body, (x1 + 30, y1 + 166), 23, COL["muted"], bold=True, max_w=x2 - x1 - 60, line_gap=6)


def bullet(draw, text, x, y, accent=COL["lime"], max_w=620):
    draw.ellipse((x, y + 10, x + 14, y + 24), fill=accent)
    draw_text(draw, text, (x + 32, y), 24, COL["white"], bold=True, max_w=max_w, line_gap=5)


def heatmap(draw, x, y, w, h):
    rounded(draw, (x, y, x + w, y + h), COL["pitch"], outline="#69B36C", width=3, r=18)
    labels = ["Short", "Own half", "Midfield", "Opp half", "Long"]
    vals = [["80%", "62%", "71%", "42%", "78%"], ["67%", "51%", "58%", "33%", "70%"]]
    fills = [[COL["lime"], COL["amber"], COL["lime"], COL["red"], COL["lime"]],
             [COL["lime"], COL["amber"], COL["amber"], COL["red"], COL["lime"]]]
    cw, ch = w // 5, h // 2
    for r in range(2):
        for c in range(5):
            cell = (x + c * cw, y + r * ch, x + (c + 1) * cw, y + (r + 1) * ch)
            draw.rectangle(cell, fill=fills[r][c], outline=COL["bg"], width=2)
            tw, th = text_size(draw, vals[r][c], font(32, bold=True))
            draw.text((cell[0] + (cw - tw) / 2, cell[1] + (ch - th) / 2), vals[r][c], font=font(32, bold=True), fill=COL["bg"])
    draw.line((x + w // 2, y, x + w // 2, y + h), fill=COL["white"], width=3)
    draw.line((x, y + h // 2, x + w, y + h // 2), fill=COL["white"], width=3)
    for i, label in enumerate(labels):
        tw, _ = text_size(draw, label, font(18, bold=True))
        draw.text((x + i * cw + (cw - tw) / 2, y + h + 24), label, font=font(18, bold=True), fill=COL["muted"])


def slide_1():
    img, draw = base_slide()
    logo(draw)
    draw_text(draw, "VOICE-FIRST HURLING ANALYTICS", (108, 214), 20, COL["orange"], bold=True)
    draw_text(draw, "Sideline AI for\nhurling coaches", (104, 278), 72, COL["white"], black=True, line_gap=14)
    draw_text(draw, "Speak the stat. Confirm it. GAAstat logs it into live match analytics.", (110, 486), 33, COL["muted"], bold=True, max_w=760)
    pill(draw, "90 second pitch", (110, 630, 330, 682), COL["lime"])
    pill(draw, "Online AI + offline manual fallback", (356, 630, 760, 682), COL["orange"])
    draw_text(draw, "Current app truth: manual match logging is offline-first; Sideline AI needs a live OpenAI Realtime connection.", (110, 960), 22, COL["muted"], bold=True, max_w=920)
    phone(draw, 1240, 112, 0.88)
    return img


def slide_2():
    img, draw = base_slide()
    header(draw, "01 / Problem")
    title_block(draw, "Match-day reality", "The coach has no spare hands", "Hurling moves too fast for paper notes and too physical for deep phone menus.")
    card(draw, (108, 560, 568, 850), "Speed", "Scores, wides, puckouts, turnovers and subs can arrive in the same passage of play.", COL["lime"])
    card(draw, (730, 560, 1190, 850), "Signal", "Many grounds have poor coverage. Cloud-only tools fail at the exact wrong time.", COL["orange"])
    card(draw, (1352, 560, 1812, 850), "Afterwards", "Coaches still need structured accountability after the final whistle.", COL["blue"])
    rounded(draw, (260, 930, 1660, 1000), COL["panel2"], outline=COL["stroke"], width=2, r=20)
    draw_text(draw, "The gap: live decisions are voice-speed, but most data capture is hand-speed.", (330, 950), 30, COL["lime"], bold=True)
    return img


def slide_3():
    img, draw = base_slide()
    header(draw, "02 / Product")
    title_block(draw, "Voice-first workflow", "Voice stats in four steps", "Sideline AI uses app tools. It is not loose chat.")
    steps = [
        ("1", "Speak", '"Point for number 11"'),
        ("2", "Confirm", '"Pending point #11. Confirm?"'),
        ("3", "Log", "Writes to match state"),
        ("4", "Ask", "Score, leaders, player stats, puckouts"),
    ]
    x = 110
    for n, title, body in steps:
        rounded(draw, (x, 520, x + 385, 740), COL["panel"], outline=COL["stroke"], width=2, r=22)
        draw.ellipse((x + 26, 548, x + 76, 598), fill=COL["lime"])
        draw_text(draw, n, (x + 42, 557), 24, COL["bg"], bold=True)
        draw_text(draw, title, (x + 28, 628), 32, COL["white"], black=True)
        draw_text(draw, body, (x + 28, 678), 22, COL["muted"], bold=True, max_w=320)
        x += 445
    rounded(draw, (140, 820, 895, 952), COL["panel2"], outline=COL["stroke"], width=2, r=20)
    draw_text(draw, "Voice writes supported now", (178, 848), 22, COL["lime"], bold=True)
    draw_text(draw, "Goals, points, wides, player/custom stats, puckouts by zone, opposition scores and undo.", (178, 888), 23, COL["white"], bold=True, max_w=650)
    rounded(draw, (1025, 820, 1780, 952), COL["panel2"], outline=COL["stroke"], width=2, r=20)
    draw_text(draw, "Not voice-supported yet", (1063, 848), 22, COL["orange"], bold=True)
    draw_text(draw, "Substitutions, notes and sync changes remain manual in the current assistant prompt.", (1063, 888), 23, COL["white"], bold=True, max_w=650)
    return img


def slide_4():
    img, draw = base_slide()
    header(draw, "03 / Reliability")
    draw_text(draw, "THE HONEST OFFLINE STORY", (100, 220), 20, COL["orange"], bold=True)
    draw_text(draw, "Offline core.\nOnline voice.", (100, 290), 68, COL["white"], black=True, line_gap=12)
    draw_text(draw, "This is the pitch claim that matches the current implementation.", (104, 480), 29, COL["muted"], bold=True, max_w=880)
    rounded(draw, (140, 570, 890, 840), COL["panel"], outline=COL["stroke"], width=2, r=24)
    draw_text(draw, "Works with no signal", (182, 612), 38, COL["lime"], black=True)
    bullet(draw, "Manual match logging, timer and drafts", 190, 688, COL["lime"])
    bullet(draw, "IndexedDB stores squad and matches locally", 190, 746, COL["lime"])
    bullet(draw, "Outbox queues Supabase sync until online", 190, 804, COL["lime"])
    rounded(draw, (1030, 570, 1780, 840), COL["panel"], outline=COL["stroke"], width=2, r=24)
    draw_text(draw, "Needs a connection", (1072, 612), 38, COL["orange"], black=True)
    bullet(draw, "Sideline AI Realtime session and WebRTC call", 1080, 688, COL["orange"])
    bullet(draw, "Initial sign-in if session is not cached", 1080, 746, COL["orange"])
    bullet(draw, "Live sharing and cloud sync transport", 1080, 804, COL["orange"])
    rounded(draw, (240, 910, 1680, 984), COL["lime"], r=20)
    draw_text(draw, "Manual logging is the offline guarantee. Voice is the fast lane when coverage is available.", (310, 932), 28, COL["bg"], bold=True)
    return img


def slide_5():
    img, draw = base_slide()
    header(draw, "04 / Hurling data")
    title_block(draw, "Not a generic voice recorder", "A hurling-specific data model", "Voice actions feed structured match data coaches can use during and after the game.")
    heatmap(draw, 132, 540, 760, 270)
    draw_text(draw, "10-zone puckout map", (330, 858), 26, COL["lime"], bold=True)
    rounded(draw, (1060, 500, 1780, 854), COL["panel"], outline=COL["stroke"], width=2, r=24)
    bullet(draw, "Puckout outcomes by zone, player and opposition winner", 1110, 550, COL["lime"], max_w=560)
    bullet(draw, "Opposition scores linked to scorer and marker", 1110, 625, COL["orange"], max_w=560)
    bullet(draw, "Pitch locations for scores and key stats", 1110, 700, COL["blue"], max_w=560)
    bullet(draw, "Player impact, targets, reports and live viewer mode", 1110, 775, COL["lime"], max_w=560)
    return img


def slide_6():
    img, draw = base_slide()
    header(draw, "05 / Close")
    draw_text(draw, "PILOT THE VOICE-LED SIDELINE ASSISTANT", (100, 220), 20, COL["orange"], bold=True)
    draw_text(draw, "Pilot Sideline AI", (100, 300), 70, COL["white"], black=True)
    draw_text(draw, "One coach -> one club -> live shared match data", (104, 405), 34, COL["muted"], bold=True, max_w=1000)
    draw_text(draw, "Free logging builds habit. Paid plans unlock analytics, reports, collaboration and live sharing.", (104, 462), 28, COL["muted"], bold=True, max_w=1120)
    plans = [
        ("Free", "EUR 0", "Core logging"),
        ("Personal Pro", "EUR 7.99/mo", "Analytics + reports"),
        ("Club", "EUR 15/mo", "Coaches + teams"),
        ("Club Pro", "EUR 25/mo", "Live viewer mode"),
    ]
    x = 120
    for name, price, body in plans:
        rounded(draw, (x, 590, x + 390, 790), COL["panel"], outline=COL["lime"] if name == "Personal Pro" else COL["stroke"], width=4 if name == "Personal Pro" else 2, r=22)
        draw_text(draw, name, (x + 28, 628), 28, COL["white"], black=True)
        draw_text(draw, price, (x + 28, 676), 28, COL["lime"] if name == "Personal Pro" else COL["orange"], black=True)
        draw_text(draw, body, (x + 28, 732), 22, COL["muted"], bold=True)
        x += 440
    rounded(draw, (210, 870, 1710, 950), COL["lime"], r=24)
    draw_text(draw, "Ask: pilot with real coaches, measure missed events and willingness to pay.", (280, 895), 33, COL["bg"], black=True)
    draw_text(draw, "Close line: a voice-controlled analyst for match day, with manual offline logging when signal disappears.", (330, 990), 24, COL["muted"], bold=True)
    return img


SCRIPT_MD = """# GAAstat voice-first 90-second pitch script

## Slide 1 - Title (0:00-0:12)
GAAstat is a voice-first hurling analytics app for match day. The main feature is Sideline AI: a coach speaks the stat, confirms it, and the app logs it straight into the live match record.

## Slide 2 - Problem (0:12-0:27)
The problem is simple. Hurling is too fast for paper notes and too physical for deep phone menus. Scores, wides, puckouts, turnovers and substitutions can happen inside one passage of play, and coaches still need accountability after the final whistle.

## Slide 3 - Product (0:27-0:48)
With Sideline AI, the workflow is speak, confirm, log and analyse. The current app can voice-log goals, points, wides, player stats, custom stats, puckouts by zone, opposition scores and undo actions. It can also answer live questions about the score, clock, player stats, leaders and puckouts.

## Slide 4 - Reliability (0:48-1:05)
The offline story is important. The core app is offline-first: manual logging, timer, drafts and local match storage work on the device, with a sync outbox for later. Sideline AI itself is online-only because it uses OpenAI Realtime over WebRTC.

## Slide 5 - Hurling Data (1:05-1:20)
This is not a generic voice recorder. The voice actions feed a hurling-specific data model: 10-zone puckout analysis, opposition scorer and marker tracking, pitch locations, player impact, targets, reports and live viewer mode for selectors.

## Slide 6 - Close (1:20-1:30)
The ask is to pilot it with real coaches this season, measure whether voice capture reduces missed events, and convert the teams that need reliable match-day analytics every week.
"""


AUDIT_MD = """# GAAstat current app audit for pitch claims

## Main pitch direction

Lead with Sideline AI, not generic offline stats:

GAAstat is a voice-first hurling match analytics app. Coaches can speak match events into Sideline AI, confirm each write, and have those events land in structured live match data. Manual logging remains the offline fallback.

## Current implemented app areas

- Live match logging: points, goals, wides, tackles, blocks, turnovers, frees, custom stats, substitutions, notes, pitch locations, opposition scores and puckouts.
- Sideline AI: embedded in the match screen, uses OpenAI Realtime through server-side SvelteKit routes, supports hold-to-talk, transcript display, pending action confirmation and local fallback parsing.
- Sideline AI write commands currently supported: goals, points, wides, add/remove player stats including custom stats, puckout won/lost with exact zone, opposition goals/points and undo last event.
- Sideline AI read commands currently supported: match summary, score, player stat leaders, individual player quick stats, team stat totals, puckout summary, recent events and current period/time.
- Do not claim voice support for substitutions, notes or sync changes yet.
- Hurling-specific analytics: 10-zone puckout heatmap, puckout by player/opposition winner, scores conceded by marker and opposition player, pitch maps, timelines, player stats, team stats, targets, coaching insights, work-ons and PDF reports.
- Data/sync: IndexedDB stores squad, matches, drafts, device state and a durable sync outbox. Supabase sync drains when online and uses backoff on failures.
- Club features: team/club support, join codes, live match sharing and live viewer mode for Club Pro.

## Offline truth

- Accurate claim: manual match logging is offline-first after the app has loaded/cached and the user is signed in.
- Accurate claim: drafts, timer, local match data and the sync outbox are local-device first.
- Accurate claim: Supabase cloud backup waits until the device is online.
- Accurate limitation: Sideline AI does not work offline in the current implementation.
- Why: the current voice feature calls `/api/realtime/call`, which sends a WebRTC offer to `https://api.openai.com/v1/realtime/calls`, and also has a session route for OpenAI Realtime client secrets. That requires network and `OPENAI_API_KEY`.
- Accurate limitation: live sharing also needs Supabase network access.
- Avoid saying: "Every stat, including voice, works offline." That is false.
"""


def render_slides():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    slide_fns = [slide_1, slide_2, slide_3, slide_4, slide_5, slide_6]
    paths = []
    for i, fn in enumerate(slide_fns, start=1):
        path = SLIDE_DIR / f"slide_{i:02d}.png"
        fn().save(path)
        paths.append(path)
    return paths


def build_pptx(paths, dest):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(dest)


def main():
    paths = render_slides()
    build_pptx(paths, PPTX_PATH)
    build_pptx(paths, VOICE_PPTX_PATH)
    SCRIPT_PATH.write_text(SCRIPT_MD, encoding="utf-8")
    AUDIT_PATH.write_text(AUDIT_MD, encoding="utf-8")
    print(PPTX_PATH)
    print(VOICE_PPTX_PATH)
    print(SCRIPT_PATH)
    print(AUDIT_PATH)


if __name__ == "__main__":
    main()
